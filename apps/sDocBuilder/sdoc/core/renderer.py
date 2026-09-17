# -*- coding: utf-8 -*-
"""sDocBuilder 渲染器：大纲数据 -> 预览统计/HTML / 技术方案 docx。

排版规格集中在 LAYOUT 表（对照 CS 人工基准模板抽取）：
A4、上下1.8/左2.5/右2.0 边距；H1 黑体16 / H2 黑体14 / H3 黑体12 / 块小节
黑体12 / 正文宋体12 行距1.3；表格素面细线；封面字号体系（客户24/文档28/
单位日期18）；目录 TOC 域（Word COM 打开时 Fields.Update 自动生成）；
页眉：COPA B=公司全称文字、COPA A=awlogo 图片；页脚居中页码域。
块内容模型见 library_store.BlockDef.items（text/image/table 三类）。
"""
import os
import re
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from sdoc.core.library_store import KIND_LABEL

PREVIEW_IMG_W = 380          # 候选池预览窗图片显示宽(px)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ====排版规格（CS 人工基准模板抽取）====
LAYOUT = {
    "page": dict(w=21.0, h=29.7, top=1.8, bottom=1.8, left=2.5, right=2.0),
    "h1": ("黑体", 16), "h2": ("黑体", 14), "h3": ("黑体", 12),
    "block_head": ("黑体", 12),
    "body": ("宋体", 12), "body_line": 1.3, "body_after": 6,
    "table": ("宋体", 10.5),
    "cover_customer": ("黑体", 24), "cover_doc": ("黑体", 28),
    "cover_info": ("宋体", 18), "cover_title": ("黑体", 26),
    "toc_head": ("黑体", 16),
    "img_max_w": 14.0, "img_max_h": 16.0,
    "header_text": ("宋体", 9), "header_logo_h": 1.1,
}
BRAND_FULL = {"COPA B": "COPA B Co., Ltd.",
              "COPA A": "COPA A LLC."}
EMBEDABLE_EXTS = (".png", ".jpg", ".jpeg")


def _set_cn_font(run, name="宋体", size=12):
    run.font.name = name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def _body_para(doc):
    """正文段落基线：行距/段后"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = LAYOUT["body_line"]
    p.paragraph_format.space_after = Pt(LAYOUT["body_after"])
    return p


def _fld(run, instr, hint=""):
    """在 run 内插入简单域（页码/TOC 等）"""
    r = run._r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = hint
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r.extend([b, it, s, t, e])


def numbered_nodes(items):
    """统一流水编号遍历：章节与块都占号（1 / 1.1 / 1.1.1，与 UI renumber 同规则）
    -> [(code, title, kind, node)]"""
    out = []

    def walk(children, prefix):
        for i, n in enumerate(children, 1):
            code = "{}.{}".format(prefix, i) if prefix else str(i)
            out.append((code, n["title"], n["kind"], n))
            if n["kind"] == "chapter":
                walk(n.get("children", []), code)

    walk(items, "")
    return out


def _filtered_table(table, brand=""):
    """按「归属公司」标签处理表格（渲染视图，不改库数据）：
    - 归属列永远不进文档（删列）；
    - 有抬头时只留 匹配抬头 或 无归属 的行；无抬头保留全部行；
    - 「序号」列过滤后重排 1..n。"""
    rows = table.get("rows") or []
    flt = table.get("filter")
    if not rows or not flt or flt not in rows[0]:
        return table
    header = rows[0]
    ci = header.index(flt)
    body = rows[1:]
    if brand:
        body = [r for r in body
                if len(r) <= ci or r[ci].strip() in ("", brand)]
    si = header.index("序号") if "序号" in header else -1
    if si > ci:
        si -= 1                          # 列删除后序号列左移
    out = [[c for j, c in enumerate(r) if j != ci] for r in [header] + body]
    if si >= 0:
        for n, r in enumerate(out[1:], 1):
            if si < len(r):
                r[si] = str(n)
    return {"rows": out, "caption": table.get("caption", "")}


class Renderer:
    def __init__(self, store):
        self.store = store
        self.missing = []          # 本次输出的占位汇总 [(位置, 说明)]
        self._project = None
        self._brand = ""           # 公司抬头（业绩表按其筛行并删归属列）

    def _entry(self, quote_id):
        for e in (self._project.entries if self._project else []):
            if e.quote_id == quote_id:
                return e
        return None

    # ----统计与摘要----

    def preview_stats(self, items, project):
        numbered = numbered_nodes(items)
        n_ch = sum(1 for _, _, k, _ in numbered if k == "chapter")
        n_blocks = imgs = missing = tables = 0
        for _, _, k, node in numbered:
            if k != "block":
                continue
            b = self.store.block(node["ref"])
            if b is None:
                continue
            n_blocks += 1
            tables += b.n_tables()
            imgs += b.n_images
            missing += b.n_missing
        return ("项目：{}　客户：{}　章节：{}　内容块：{}　表格：{}　"
                "图片：{}张（缺 {}）".format(
                    project.project_id, project.company, n_ch,
                    n_blocks, tables, imgs, missing))

    def preview_html(self, items, project):
        rows = []
        for num, title, kind, node in numbered_nodes(items):
            if kind == "chapter":
                level = min(num.count(".") + 2, 4)
                rows.append("<h{}>{} {}</h{}>".format(level, num, title,
                                                      level))
                continue
            b = self.store.block(node["ref"])
            if b is None:
                rows.append("<p style='color:#b00'>【未知块：{}】</p>"
                            .format(node["ref"]))
                continue
            note = []
            if b.n_paras():
                note.append("{} 段".format(b.n_paras()))
            if b.n_tables():
                note.append("{} 表".format(b.n_tables()))
            if b.n_images or b.n_missing:
                note.append("{} 图（缺 {}）".format(b.n_images,
                                                  b.n_missing))
            rows.append("<p style='margin-left:2em;color:#555'><b>{} {}</b>"
                        " ［{}］—— {}</p>".format(
                            num, title, KIND_LABEL.get(b.kind, b.kind),
                            "、".join(note)))
        head = "<div><b>{}</b></div><hr/>".format(
            self.preview_stats(items, project))
        return head + "\n".join(rows)

    # ----候选池点击预览（HTML）----

    def blocks_html(self, refs, project=None, brand=""):
        """refs（块或 entry 配置）-> 预览 HTML"""
        if project is not None and self._project is not project:
            self._project = project
        self._brand = brand
        parts = ["<div style='color:#888'>材料预览</div><hr/>"]
        for ref in refs:
            if ref.startswith("entry:"):
                e = self._entry(int(ref.split(":")[1])
                                if ref.split(":")[1].isdigit() else -1)
                if e is None:
                    continue
                parts.append(self._entry_config_html(e))
                continue
            b = self.store.block(ref)
            if b is None:
                continue
            parts.append("<h3>{}</h3>".format(b.title))
            for kind, payload in b.items:
                if kind == "text":
                    t = payload
                    if t.startswith("## "):
                        parts.append("<p><b>{}</b></p>".format(t[3:].strip()))
                    else:
                        parts.append("<p>{}</p>".format(t))
                elif kind == "image":
                    path, cap = payload
                    if os.path.exists(path):
                        parts.append("<p><img src='{}' width='{}'><br/>"
                                     "<span style='color:#888'>{}</span></p>"
                                     .format(_file_url(path),
                                             PREVIEW_IMG_W,
                                             cap or ""))
                    else:
                        parts.append("<p style='color:#b00'>【待补充：{}】"
                                     "</p>".format(cap or os.path.basename(path)))
                elif kind == "table":
                    parts.append(_rows_html(
                        _filtered_table(payload,
                                        brand).get("rows") or [],
                        payload.get("caption") or ""))
        return "".join(parts)

    def _entry_config_html(self, e):
        sel = e.selection or {}
        hd = sel.get("header") or {}
        md = e.engine.get("module_data") or {}
        cfg = [["项目", "内容"]]
        if sel.get("module"):
            cfg.append(["称重模块（含传感器）", " ".join(
                x for x in (sel["module"] if isinstance(sel["module"], list)
                            else [sel["module"]]))])
        if sel.get("platform"):
            cfg.append(["平台秤秤体", sel["platform"]])
        if sel.get("bench"):
            cfg.append(["台秤秤体", sel["bench"]])
        if sel.get("jbox"):
            cfg.append(["接线盒", sel["jbox"]])
        if sel.get("controller"):
            cfg.append(["显示仪表", "{}（{}）".format(
                sel["controller"],
                (sel.get("controller_install") or "").strip("，, ")
                or "安装方式见技术说明")])
        acc = sel.get("accessories") or []
        if acc:
            cfg.append(["选配附件", "、".join(acc)])
        spec = [["参数", "值"]]
        for label, key in (("最大量程", "range"), ("检定分度值", "e")):
            if hd.get(key) not in (None, "", "0"):
                spec.append([label, str(hd[key])])
        if hd.get("com2") not in (None, "", "0"):
            spec.append(["通讯", str(hd["com2"])])
        if hd.get("power"):
            spec.append(["工作电源", str(hd["power"])])
        spec.append(["防爆要求", "是（{}）".format(
            e.module_data("required_ex") or "等级详见选型") if e.ex else "否"])
        if e.digital:
            spec.append(["数字化方案", "是"])
        if e.metrology:
            spec.append(["计量检定", "需要检定"])
        for label, key in (("材质", "material"), ("支点数", "support"),
                           ("台面尺寸", "台面尺寸")):
            v = str(md.get(key, "") or "").strip()
            if v and v not in ("（未提供）", "（空）", "None"):
                spec.append([label, v])
        html = "<h3>{}（快照条目 {}）</h3>".format(e.type_cn, e.quote_id)
        html += _rows_html(cfg, "")
        html += "<p><b>技术参数</b></p>" + _rows_html(spec, "")
        if e.description:
            html += "<p>需求描述：{}</p>".format(e.description)
        return html

    # ----docx 输出----

    def _setup_page(self, doc):
        s = doc.sections[0]
        pg = LAYOUT["page"]
        s.page_width, s.page_height = Cm(pg["w"]), Cm(pg["h"])
        s.top_margin, s.bottom_margin = Cm(pg["top"]), Cm(pg["bottom"])
        s.left_margin, s.right_margin = Cm(pg["left"]), Cm(pg["right"])

    def _setup_header_footer(self, doc, brand):
        """正文节页眉（COPA B=文字 / COPA A=logo）与页脚页码；封面节留空"""
        body = doc.sections[1]
        body.header.is_linked_to_previous = False
        body.footer.is_linked_to_previous = False
        hp = body.header.paragraphs[0]
        if brand == "COPA B":
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            name, size = LAYOUT["header_text"]
            _set_cn_font(hp.add_run(
                BRAND_FULL["COPA B"] + "技术文件"), name, size)
        elif brand == "COPA A":
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT   # logo 页眉右侧
            logo = self._aw_logo()
            if logo:
                hp.add_run().add_picture(logo, height=Cm(
                    LAYOUT["header_logo_h"]))
        fp = body.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        _set_cn_font(run, "宋体", 9)
        _fld(run, " PAGE ", "1")

    def _aw_logo(self):
        root = getattr(self.store, "root", None)
        if not root:
            return None
        d = os.path.join(root, "_global", "awlogo")
        if not os.path.isdir(d):
            return None
        for name in sorted(os.listdir(d)):
            if name.lower().endswith(EMBEDABLE_EXTS):
                return os.path.join(d, name)
        return None

    def _cover(self, doc, project, brand):
        """封面（CS 基准字号体系；本节垂直居中由 export 统一设）。

        版序：客户名称 -> 项目名称 -> 技术方案(大标题) -> 品牌全称 ->
        日期；项目名称与客户名称同一字体字号(黑体24)。"""
        for text, style_key in (
                (project.company, "cover_customer"),
                (project.project_id, "cover_customer"),
                ("技术方案", "cover_doc"),
                (BRAND_FULL.get(brand, ""), "cover_info"),
                (datetime.now().strftime("%Y年%m月%d日"), "cover_info")):
            if not text:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            name, size = LAYOUT[style_key]
            _set_cn_font(p.add_run(text), name, size)

    def _toc(self, doc):
        """目录标题 + TOC 域（Word 打开/预览转换时 Fields.Update 生成）"""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name, size = LAYOUT["toc_head"]
        _set_cn_font(p.add_run("目  录"), name, size)
        run = doc.add_paragraph().add_run()
        _set_cn_font(run, "宋体", 12)
        _fld(run, ' TOC \\o "1-3" \\h \\z \\u ',
             "（目录将在 Word 打开时自动更新）")

    def export_docx(self, items, project, out_dir, out_path=None, brand=""):
        self.missing = []
        self._project = project
        self._brand = brand
        doc = Document()
        self._setup_page(doc)
        self._cover(doc, project, brand)
        doc.add_section(WD_SECTION.NEW_PAGE)
        from docx.oxml import OxmlElement

        for sect, val in ((doc.sections[0], "center"),
                          (doc.sections[1], "top")):
            vAlign = OxmlElement("w:vAlign")
            vAlign.set(qn("w:val"), val)
            sect._sectPr.append(vAlign)
        self._setup_header_footer(doc, brand)
        self._toc(doc)
        doc.add_page_break()

        sub = [0]                      # 当前章的语料小节序号（跨块连续）
        chapter = [""]                 # 最近章编号（语料小节挂它：3.1.1…）
        for num, title, kind, node in numbered_nodes(items):
            if kind == "chapter":
                sub[0] = 0
                chapter[0] = num
                level = min(num.count(".") + 1, 3)
                h = doc.add_heading("", level=level)
                name, size = LAYOUT["h{}".format(level)]
                run = h.add_run("{} {}".format(num, title))
                _set_cn_font(run, name, size)
                run.font.color.rgb = RGBColor(0, 0, 0)   # 标题一律黑色
            else:
                # 块标题行不占编号；统一按「隐形块名」输出：1pt 白字 +
                # 固定10磅行距——块名保留在文档中（可寻址、后期改色即可
                # 恢复显示），版面上完全不可见
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = \
                    WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(10)
                run = p.add_run(title)
                _set_cn_font(run, "黑体", 1)
                run.font.color.rgb = RGBColor(255, 255, 255)
                self._emit_block(doc, node["ref"], title,
                                 "{} {}".format(num, title),
                                 chapter_code=chapter[0], sub=sub)
        os.makedirs(out_dir, exist_ok=True)
        path = out_path or os.path.join(
            out_dir, "技术方案_{}_{}.docx".format(
                project.project_id, datetime.now().strftime("%Y%m%d_%H%M%S")))
        doc.save(path)
        return path

    # ----块内容发射----

    def _emit_block(self, doc, ref, btitle, where,
                    chapter_code="", sub=None):
        """块内容发射。语料内部小节（## 子标题、图前题注短行）按
        「章编号.序号」梯度编号（3.1.1 产品描述 / 3.1.2 技术指标…），
        序号跨块连续（sub 由 export 按章维护）；块标题行本身不占号。"""
        if ref.startswith("entry:"):       # 条目选型配置（来自快照，不入库）
            self._emit_entry_config(doc, ref.split(":")[1],
                                    chapter_code, sub)
            return
        b = self.store.block(ref)
        if b is None:
            _body_para(doc).add_run("【未知块：{}】".format(ref))
            return
        if sub is None:
            sub = [0]

        def _next_sub():
            sub[0] += 1
            return "{}.{}".format(chapter_code, sub[0]) if chapter_code \
                else ""

        for idx, (kind, payload) in enumerate(b.items):
            if kind == "text":
                text = payload
                if text.startswith("## "):
                    p = doc.add_paragraph()
                    label = text[3:].strip()
                    code = _next_sub()
                    if code:
                        label = "{} {}".format(code, label)
                    _set_cn_font(p.add_run(label), "黑体", 12)
                else:
                    nxt = b.items[idx + 1][0] \
                        if idx + 1 < len(b.items) else None
                    if nxt == "image" and len(text) <= 20 and chapter_code:
                        # 图前题注短行（如「技术指标」）也占小节号
                        p = doc.add_paragraph()
                        _set_cn_font(p.add_run("{} {}".format(
                            _next_sub(), text)), "黑体", 12)
                    else:
                        _set_cn_font(_body_para(doc).add_run(text),
                                     *LAYOUT["body"])
            elif kind == "image":
                path, cap = payload
                if os.path.exists(path) \
                        and path.lower().endswith(EMBEDABLE_EXTS):
                    try:
                        w = self._fit_image(path)
                        doc.add_picture(path, width=Cm(w))
                        doc.paragraphs[-1].alignment = \
                            WD_ALIGN_PARAGRAPH.CENTER
                    except Exception:
                        self._placeholder(doc, where,
                                          cap or os.path.basename(path))
                else:
                    self._placeholder(doc, where, cap or os.path.basename(path))
            elif kind == "table":
                self._emit_table(doc, _filtered_table(payload, self._brand))
        _body_para(doc)

    def _fit_image(self, path):
        """自适应显示宽：原始尺寸/DPI 推自然宽，只缩不放；宽高双限
        （修超长图撑页：高超限时按高反缩）"""
        from PIL import Image
        with Image.open(path) as im:
            px_w, px_h = im.size
            try:
                dpi = im.info.get("dpi", (96, 96))[0] or 96
            except Exception:
                dpi = 96
        nat_w = px_w / float(dpi) * 2.54
        nat_h = px_h / float(dpi) * 2.54
        w = min(nat_w, LAYOUT["img_max_w"])
        if w <= 0:
            return LAYOUT["img_max_w"]
        if w * nat_h / nat_w > LAYOUT["img_max_h"]:
            w = LAYOUT["img_max_h"] * nat_w / nat_h
        return round(w, 2)

    def _emit_entry_config(self, doc, quote_id, chapter_code="", sub=None):
        """条目选型配置（快照参数 -> 技术文档所需的少量参数）；
        内部小标题占语料小节号"""
        e = self._entry(int(quote_id) if quote_id.isdigit() else -1)
        if e is None:
            _body_para(doc).add_run("【条目 {} 未加载】".format(quote_id))
            return
        if sub is None:
            sub = [0]

        def _next_sub():
            sub[0] += 1
            return "{}.{}".format(chapter_code, sub[0]) if chapter_code \
                else ""

        sel = e.selection or {}
        hd = sel.get("header") or {}
        md = e.engine.get("module_data") or {}

        p = doc.add_paragraph()
        _set_cn_font(p.add_run("{} 选型配置".format(_next_sub())),
                     "黑体", 12)
        rows = [["项目", "内容"]]
        if sel.get("module"):
            rows.append(["称重模块（含传感器）", " ".join(
                x for x in (sel["module"] if isinstance(sel["module"], list)
                            else [sel["module"]]))])
        if sel.get("platform"):
            rows.append(["平台秤秤体", sel["platform"]])
        if sel.get("bench"):
            rows.append(["台秤秤体", sel["bench"]])
        if sel.get("jbox"):
            rows.append(["接线盒", sel["jbox"]])
        if sel.get("controller"):
            rows.append(["显示仪表", "{}（{}）".format(
                sel["controller"],
                (sel.get("controller_install") or "").strip("，, ")
                or "安装方式见技术说明")])
        acc = sel.get("accessories") or []
        if acc:
            rows.append(["选配附件", "、".join(acc)])
        self._emit_table(doc, {"rows": rows, "caption": ""})

        p = doc.add_paragraph()
        _set_cn_font(p.add_run("{} 技术参数".format(_next_sub())),
                     "黑体", 12)
        rows = [["参数", "值"]]
        for label, key in (("最大量程", "range"), ("检定分度值", "e")):
            if hd.get(key) not in (None, "", "0"):
                rows.append([label, str(hd[key])])
        if hd.get("com2") not in (None, "", "0"):
            rows.append(["通讯", str(hd["com2"])])
        if hd.get("power"):
            rows.append(["工作电源", str(hd["power"])])
        if e.ex:
            rows.append(["防爆要求", "是（{}）".format(
                e.module_data("required_ex") or "等级详见选型")])
        else:
            rows.append(["防爆要求", "否"])
        if e.digital:
            rows.append(["数字化方案", "是"])
        if e.metrology:
            rows.append(["计量检定", "需要检定"])
        for label, key in (("材质", "material"), ("支点数", "support"),
                           ("台面尺寸", "台面尺寸")):
            v = str(md.get(key, "") or "").strip()
            if v and v not in ("（未提供）", "（空）", "None"):
                rows.append([label, v])
        self._emit_table(doc, {"rows": rows, "caption": ""})

        if e.description:
            p = _body_para(doc)
            _set_cn_font(p.add_run("需求描述：{}".format(e.description)),
                         *LAYOUT["body"])
        _body_para(doc)

    def _emit_table(self, doc, table):
        rows = table.get("rows") or []
        cap = table.get("caption") or ""
        if not rows:
            return
        if cap:
            p = doc.add_paragraph()
            _set_cn_font(p.add_run(cap), "黑体", 10)
        width = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=width)
        t.style = "Table Grid"
        self._thin_borders(t)
        for r, row in enumerate(rows):
            for c in range(width):
                cell = row[c] if c < len(row) else ""
                p = t.cell(r, c).paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                _set_cn_font(p.add_run(cell), *LAYOUT["table"])

    @staticmethod
    def _thin_borders(t):
        """素面细边框（0.5pt 灰），替换 Table Grid 默认粗黑框"""
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        tblPr.append(borders)

    def _placeholder(self, doc, where, what):
        self.missing.append((where, what))
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cn_font(p.add_run("【待补充：{}】".format(what)), "宋体", 12)


def _file_url(path):
    return "file:///" + os.path.abspath(path).replace("\\", "/")


def _rows_html(rows, caption):
    if not rows:
        return ""
    html = ["<table width='100%' cellspacing='0' cellpadding='3' "
            "style='border:1px solid #bbbbbb; font-size:12px'>"]
    if caption:
        html.append("<caption align='left'>{}</caption>".format(caption))
    for r in rows:
        html.append("<tr>" + "".join(
            "<td style='border:1px solid #bbbbbb'>{}</td>".format(
                (c or "")) for c in r) + "</tr>")
    html.append("</table>")
    return "".join(html)
