# -*- coding: utf-8 -*-
"""COPA_com_Writer
职责边界：Checklist 只负责收集人工确认后的结构化选择(型号下拉/数量/复选框)，
本模块负责：① 到 data 的 描述通表/属性通表 CSV 匹配型号调取描述
(描述统一以三张品牌描述通表为源，按型号全局搜索定位)；
② 按报价抬头选择模板并填充 xlsx(COPA B)。

填充约定(COPA B)：
- 多项目(max quote_id > 1) -> COPA B多项目报价模板.xlsx，每个 quote_id 一个分项段
  (段内：名称/型号、单位、数量、产品描述，小计、套数合计)，段后 设备合计/服务费用/总价；
- 单项目 -> COPA B单项目报价模板.xlsx 通用明细行；
- 数量：模块业务 模块=支点数、接线盒=1、套数合计=spinbox；平台秤/台秤=spinbox；
- 仪表非多通道 -> 数量1 留在项目内；多通道(可接N个秤台) -> 从项目中删去，
  在项目底部追加统一行，数量 = ceil(总套数 ÷ 通道数)；
- 输出语言：build_quote(lang="zh"/"en")，英文走 template/*_英文版 模板，
  writer 动态文案同步切换(描述通表文本仍中文，待描述翻译接入)；
  英文模板无服务费行，现场服务开关仅对中文模板生效；
  汇总/服务费/备注区行按内容锚点定位，模板增删备注行无需改代码。
"""
import csv
import math
import os
import re
import sys
import unicodedata
from copy import copy
from datetime import datetime
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# 数据根目录：冻结打包后 data/template/output 随 exe 同目录手动摆放(保持源码结构)，
# 源码运行时按 Writer 自身位置上溯定位
if getattr(sys, "frozen", False):
    _COPA_ROOT = os.path.dirname(sys.executable)
else:
    _COPA_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DATA_DIR = os.path.join(_COPA_ROOT, "data")
TEMPLATE_DIR = os.path.join(_COPA_ROOT, "template")
OUTPUT_DIR = os.path.join(_COPA_ROOT, "output")

_UNIT_NAMES = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]
_UNIT_NAMES_EN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]

# 双语单位格(参照阿根廷双语报价范例)：英文在上中文在下
_UNIT_BI = {"只": "PCS\n只", "台": "PCS\n台", "米": "M\n米",
            "块": "PCS\n块", "套": "SET\n套"}

# 输出语言资源(lang)：en 取 *_英文版 模板并切换 writer 生成的动态文案；
# 描述库文本(lookup_*/build_scale_desc 等)暂仍中文，描述翻译接入后再切换
_TEXT = {
    "zh": {
        "module": "称重模块",
        "jbox": "接线盒",
        "controller": "称重显示控制器",
        "cable": "信号电缆",
        "plate_up": "上安装过渡板",
        "plate_down": "下安装过渡板",
        "platform": "电子平台秤",
        "platform_ex": "防爆电子平台秤",
        "bench": "电子台秤",
        "display": "大屏幕",
        "display_ex": "防爆大屏幕",
        "pc": "只",
        "m": "米",
        "plate": "块",
        "scale_unit": "台",
        "set": "套",
        "sets_total": "套数合计",
        "ex_note": "防爆环境。",
    },
    "en": {
        "module": "Weighing module",
        "jbox": "Junction box",
        "controller": "Weighing display controller",
        "cable": "Signal cable",
        "plate_up": "Top adapting plate",
        "plate_down": "Base adapting plate",
        "platform": "Electronic platform scale",
        "platform_ex": "Explosion-proof electronic platform scale",
        "bench": "Electronic bench scale",
        "display": "Large display",
        "display_ex": "Explosion-proof large display",
        "pc": "pc",
        "m": "m",
        "plate": "pc",
        "scale_unit": "unit",
        "set": "set",
        "sets_total": "Sets Total",
        "ex_note": "Ex-proof environment.",
    },
}


def _txt(lang, key):
    """按输出语言取动态文案；未知语言回退中文"""
    return _TEXT.get(lang, _TEXT["zh"])[key]


def _name_text(lang, key):
    """条目名称：en 为 英文\\n中文 双语，zh 单语"""
    if lang == "en":
        return "{}\n{}".format(_TEXT["en"][key], _TEXT["zh"][key])
    return _TEXT["zh"][key]


def _unit_text(lang, key):
    """单位格：en 为 英文\\n中文 双语(参照范例 PCS/台)，zh 单语"""
    zh = _TEXT["zh"][key]
    if lang == "en" and zh in _UNIT_BI:
        return _UNIT_BI[zh]
    return _TEXT["en"][key] if lang == "en" else zh


def _desc_bi(lang, zh, en):
    """描述格：en 为 中文在上英文下(范例格式)；英文缺失回退中文"""
    if lang == "en" and (en or "").strip():
        return "{}\n{}".format(zh, en)
    return zh


def _unit_names(lang):
    """分项段序号字表：中文数字/罗马数字"""
    return _UNIT_NAMES_EN if lang == "en" else _UNIT_NAMES


def _tpl_name(base, lang):
    """模板文件名按输出语言加 _英文版 后缀"""
    if lang == "en":
        stem, ext = os.path.splitext(base)
        return stem + "_英文版" + ext

    return base

# 平台秤/台秤 CSV 材质代码 -> 中文
_MATERIAL_NAMES = {
    "C": "碳钢", "HC": "碳钢", "S": "不锈钢", "HS": "不锈钢",
    "CS": "混合材质", "HCS": "混合材质",
}


# ---------------------------------------------------------------------------
# CSV 装载与匹配
# ---------------------------------------------------------------------------

def _norm(text):
    """型号归一化：NFKC(全角→半角等) + 去空白 + 大写，用于跨表比对"""
    return unicodedata.normalize("NFKC", str(text)).replace(" ", "").upper()


def _read_csv(name):
    path = os.path.join(DATA_DIR, name)
    # utf-8-sig 兼容 Excel 另存"CSV UTF-8"时带的 BOM(首列列名不被污染)，
    # 普通 UTF-8 文件不受影响
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


# 描述索引：三张品牌描述通表(品牌,型号family,详细型号,描述)统一装载，
# 仪表/模块/手写秤体描述共用一个 型号->描述 索引(COPA B/COPA A报价均可能选用
# FT/AW/FAB 器材，型号编码不冲突)。
# 惰性加载约定(与 Parser._vocab 一致)：声明为空容器而非 None，配
# _INDEXES_BUILT 构建标记——首次用到才读 CSV，静态类型不引入 Optional
_DESC: dict = {}             # {norm型号: {"brand","family","model","desc","channels"}}
_PLATFORMS: dict = {}        # {具体型号: row}  平台秤通表.csv(FT+AW合并属性表)
_BENCHES: dict = {}          # {具体型号: row}  台秤通表.csv(FT+AW合并属性表)
_JBOXES: dict = {}           # {norm型号: row}
_VARIANT_SWAPS: dict = {}    # {衍生family: 基础family}  衍生品种映射.csv 反转
_INDEXES_BUILT = False


def _ensure_indexes():
    global _INDEXES_BUILT, _PLATFORMS, _BENCHES, _JBOXES

    if _INDEXES_BUILT:
        return

    for csv_name in ("description/FT描述通表.csv",
                     "description/AW描述通表.csv",
                     "description/FAB描述通表.csv"):
        for row in _read_csv(csv_name):
            desc = (row.get("描述") or "").strip()

            if not desc:
                continue

            m = re.search(r"可接(\d+)个秤台", desc)

            # 详细型号支持 "CS 3t|CS 4.5t|..." 管道符一行多码(共用描述)，逐码入索引
            for model in (row.get("详细型号") or "").split("|"):
                model = model.strip()

                if not model:
                    continue

                _DESC[_norm(model)] = {
                    "brand": (row.get("品牌") or "").strip(),
                    "family": (row.get("型号family") or "").strip(),
                    "model": model,
                    "desc": desc,
                    "desc_en": (row.get("描述EN") or "").strip(),
                    "channels": int(m.group(1)) if m else 1,
                }

    _PLATFORMS = {r["具体型号"]: r for r in _read_csv("平台秤通表.csv")}
    _BENCHES = {r["具体型号"]: r for r in _read_csv("台秤通表.csv")}
    _JBOXES = {_norm(r["model型号"]): r for r in _read_csv("Jbox通表 (FT+AW).csv")}

    # 衍生品种映射反转(衍生family -> 基础family)：衍生型号描述按标准型号输出的依据
    try:
        for row in _read_csv("衍生品种映射.csv"):
            to_family = (row.get("衍生family") or "").strip()
            from_family = (row.get("基础family") or "").strip()

            if to_family and from_family:
                _VARIANT_SWAPS[to_family] = from_family
    except OSError:
        pass

    _INDEXES_BUILT = True


def lookup_instrument(model):
    """仪表详细型号 -> {"model","desc","channels"}；未匹配返回 None"""
    if not model:
        return None

    _ensure_indexes()
    return _DESC.get(_norm(model))


def lookup_accessory(model):
    """附件型号(接线盒/过渡板等) -> 描述。
    统一查品牌描述通表(型号family列仅备注不参与过滤)；
    接线盒无描述行时按 Jbox 通表属性合成(不涉及线制)"""
    if not model:
        return ""

    _ensure_indexes()
    entry = _DESC.get(_norm(model))

    if entry is not None and entry["desc"]:
        return entry["desc"]

    jbox = _JBOXES.get(_norm(model))

    if jbox:
        ex = jbox["EX防爆等级"]
        brand = "flintec" if jbox["品牌"] == "FT" else "alwe"
        # 业务约定：接线盒描述不涉及线制，仅给出线数与防爆等级
        return "不锈钢接线盒，{}出线{}；品牌：{}".format(
            jbox["出线数"],
            "，防爆等级：" + ex if ex != "0" else "",
            brand,
        )

    return ""


def lookup_desc(model):
    """品牌描述通表 -> 描述(仪表/模块/电缆/过渡板等统一索引)；未匹配返回空串"""
    if not model:
        return ""

    _ensure_indexes()
    entry = _DESC.get(_norm(model))
    return entry["desc"] if entry is not None else ""


def lookup_desc_en(model):
    """品牌描述通表 -> 英文描述(描述EN 列)；未匹配/未填写返回空串"""
    if not model:
        return ""

    _ensure_indexes()
    entry = _DESC.get(_norm(model))
    return (entry.get("desc_en") or "") if entry is not None else ""


# 支架类型 -> AW 防尘式仪表尾码倒数第二位(1=立杆, 2=壁挂)与描述前缀
_BRACKET_DIGIT = {"立杆支架": "1", "壁挂支架": "2"}
_BRACKET_PREFIX = {"立杆支架": "立杆安装，", "壁挂支架": "壁挂安装，"}


def bracket_variant(model, bracket):
    """仪表型号 -> 指定支架类型的型号。AW 防尘式系列换尾码变体(改码后须存在于
    品牌通表，数据校验失败则原样返回)；FT/FAB 编码不随支架变化，原样返回。"""
    digit = _BRACKET_DIGIT.get(bracket)

    if not digit or not model.startswith("AW") or len(model) < 2:
        return model

    if not (model[-2].isdigit() and model[-2] != digit):
        return model

    candidate = model[:-2] + digit + model[-1]
    return candidate if _DESC.get(_norm(candidate)) else model


def _segment_in(token, text):
    """token 以独立段出现于 text：两侧邻字符(或文本边界)均非字母数字。
    挡住子串误命中(AW 材质"C"命中"AWC"/"C3"段、FT 型号"52-30M"命中"52-30ML"等)"""
    if not token:
        return False

    t = text.upper()
    k = token.upper()
    start = 0

    while True:
        i = t.find(k, start)

        if i < 0:
            return False

        if (i == 0 or not t[i - 1].isalnum()) \
                and (i + len(k) >= len(t) or not t[i + len(k)].isalnum()):
            return True

        start = i + 1


def lookup_module(module_model, material, family, capacity, ex, y=None,
                  accuracy=None):
    """模块+传感器条件 -> (完整规格型号, 描述)。

    品牌描述通表中的模块行(如 "52-30M CS SB14-4536kg-BH-C3 EX Y=11500")以
    段精确方式匹配：模块型号锚定行首且紧随字符非字母数字(挡住 L 变体行
    "52-30ML"/"M20L" 的子串误命中)、材质以独立段出现、传感器family-容量、
    防爆标记须同时一致(FT表为" EX" token，AW表为详细型号"-X"尾缀)；
    容量token先按kg原值(FT表)，未命中再按吨换算(AW表如 "M50-C-AWD-5t-C3"；
    库内一律kg标识，换算口径同engine.format_capacity：整千取整否则保留小数)；
    给定 accuracy(C3/C6等)时优先返回含该精度段的行，无对应行回退不分精度；
    给定 y 时优先返回 " Y={y}" 行(如 "Y=23000")，表内无对应Y行则回退首项匹配。
    """
    _ensure_indexes()

    if capacity is not None:
        cap = "%.12g" % float(capacity)
        tons = float(capacity) / 1000
        cap_tokens = [
            f"{family}-{cap}kg",
            "{}-{}t".format(family, int(tons) if tons.is_integer() else tons),
        ]
    else:
        cap_tokens = [family]

    m_up = module_model.upper()

    for cap_token in cap_tokens:
        cap_up = cap_token.upper()
        y_up = " Y={}".format(y).upper() if y not in (None, "") else None
        hits = []

        for entry in _DESC.values():
            model = entry["model"]
            up = model.upper()

            # 防爆标记双约定：FT表型号串含" EX" token；AW表详细型号以"-X"尾缀
            ex_in_model = " EX" in up or up.endswith("-X")

            # 模块型号锚定行首且紧随字符非字母数字(L变体行不误命中)；
            # 材质按独立段匹配(AW单字母材质不误命中AWC/C3等片段)
            if not (up.startswith(m_up)
                    and not up[len(m_up):len(m_up) + 1].isalnum()
                    and _segment_in(material, model)
                    and cap_up in up and (ex_in_model == bool(ex))):
                continue

            hits.append((model, entry["desc"]))

        if not hits:
            continue

        if accuracy:
            acc_hits = [h for h in hits if _segment_in(accuracy, h[0])]

            if acc_hits:
                hits = acc_hits

        if y_up is None:
            return hits[0]

        for hit_model, hit_desc in hits:
            if y_up in hit_model.upper():
                return hit_model, hit_desc

        # 表内无对应Y行 -> 回退本容量token的首项命中(同原fallback语义)
        return hits[0]

    return module_model, ""


def _scale_desc(row, unit_col):
    """平台秤/台秤属性表无描述列 -> 由属性合成描述"""
    material = _MATERIAL_NAMES.get(row["材质"], row["材质"])
    brand = "flintec" if row.get("品牌") == "FT" else "alwe"
    desc = "{}，量程{}kg，台面尺寸{}mm，{}，分度值{}{}".format(
        row["细分品种"], row["kg量程"], row["台面尺寸"],
        material, row[unit_col], "g" if unit_col == "g分度值" else "kg",
    )

    if row["防爆"] == "1":
        desc += "，防爆"

    return desc + "；品牌：" + brand


def _standard_model(model):
    """衍生型号 -> 标准基础型号(family段按 衍生品种映射.csv 反换)；非衍生原样返回"""
    parts = model.split("-")
    base_family = _VARIANT_SWAPS.get(parts[0])

    if not base_family:
        return model

    return "-".join([base_family] + parts[1:])


def _lookup_scale_desc(model, table, unit_col):
    """秤体描述优先级：品牌通表手写行 -> 属性通表合成 -> 衍生型号按标准型号输出。
    反换后的标准family不再是衍生码，递归至多一层。"""
    entry = _DESC.get(_norm(model))

    if entry is not None:
        return entry["desc"]

    row = table.get(model)

    if row is not None:
        return _scale_desc(row, unit_col)

    std = _standard_model(model)

    if std != model:
        return _lookup_scale_desc(std, table, unit_col)

    return ""


def lookup_platform(model):
    _ensure_indexes()
    return _lookup_scale_desc(model, _PLATFORMS, "kg分度值")


def lookup_bench(model):
    _ensure_indexes()
    return _lookup_scale_desc(model, _BENCHES, "g分度值")


def lookup_platform_row(model):
    """平台秤通表原始行(型批型号/材质/量程/台面/分度值/品牌/可用传感器)"""
    _ensure_indexes()
    return _PLATFORMS.get(model)


def lookup_bench_row(model):
    """台秤通表原始行"""
    _ensure_indexes()
    return _BENCHES.get(model)


# 传感器材质索引：family+容量 -> 材质原文(整机描述"（材质：..）"用)
_SENSOR_MTL: dict = {}


def _load_sensor_mtl():
    for csv_name in ("富林泰克传感器.csv", "COPA A传感器.csv"):
        try:
            rows = _read_csv(csv_name)
        except OSError:
            continue

        for row in rows:
            _SENSOR_MTL["{}|{}".format(
                (row.get("family型号") or "").strip().upper(),
                (row.get("capacity容量") or "").strip(),
            )] = (row.get("Mtl材质") or "").strip()


def sensor_material(sensor_text):
    """'SLB-227kg' 式传感器型号 -> 材质中文；未匹配返回''"""
    if not _SENSOR_MTL:
        _load_sensor_mtl()

    hit = re.match(r"([A-Za-z0-9]+)-([\d.]+)", (sensor_text or "").strip())

    if hit is None:
        return ""

    raw = _SENSOR_MTL.get("{}|{}".format(
        hit.group(1).upper(), hit.group(2))) or ""

    if "17-4PH" in raw or "不锈钢" in raw or raw in ("SS", "S"):
        return "不锈钢"

    if "合金钢" in raw:
        return "合金钢"

    return raw


# 品牌 -> 整机描述里的品牌名
_BRAND_NAMES = {"FT": "flintec", "AW": "alwe", "FAB": "flintab"}

# 通表材质代码 -> 整机描述"1.材质："文案(平台秤/台秤共用)
_SCALE_MTL_TEXT = {
    "C": "碳钢，花纹钢台面", "HC": "碳钢，花纹钢台面",
    "S": "全304不锈钢，拉丝面罩", "HS": "全304不锈钢，拉丝面罩",
    "CS": "碳钢框架，不锈钢台面", "HCS": "碳钢框架，不锈钢台面",
}

# 上表 -> 英文(整机描述英文块用，措辞对齐阿根廷双语范例)
_SCALE_MTL_TEXT_EN = {
    "C": "carbon steel, tread plate deck", "HC": "carbon steel, tread plate deck",
    "S": "full SS304 with brushed front panel", "HS": "full SS304 with brushed front panel",
    "CS": "carbon steel frame with stainless steel deck",
    "HCS": "carbon steel frame with stainless steel deck",
}


def instrument_display_prefix(model):
    """仪表详细型号 -> 展示前缀(FAB3306HX1AG00A->FAB3306，FT230HA012A->FT-230)。
    规则：品牌通表family的归一化(去-)是型号前缀时用family(带-)；
    型号在family后紧跟数字则并入该数字(3306系列的6)。"""
    entry = lookup_instrument(model)

    if entry is None:
        return str(model or "").split("(")[0]

    family = entry.get("family") or ""
    norm = family.replace("-", "").upper()
    model_text = str(model or "")

    if norm and model_text.upper().startswith(norm):
        rest = model_text[len(norm):]

        if rest[:1].isdigit():
            return norm + rest[:1]

        return family

    return family or model_text


def _scale_model_text(item):
    """'SCS-0.3t（FTPD-S-003-0606）'式型号文本：型批型号（具体型号）。
    通表无行(手输型号)时退化为具体型号本身。"""
    p_type = item["type"]
    row = (lookup_platform_row(item.get("platform"))
           if p_type == "platform" else lookup_bench_row(item.get("bench")))
    model = item.get("platform") or item.get("bench") or ""

    if row is None:
        return model

    return "{}（{}）".format(row.get("型批型号") or "", model)


def build_scale_desc(item, with_model=True, lang="zh"):
    """平台秤/台秤整机多行描述(对齐真实报价单的编号式写法)。

    with_model=True(COPA A多项目/内容块模板)：首行"型号：型批型号（具体型号）"；
    with_model=False(COPA B多项目)：型号并入名称列，描述自 1.材质 起。

    平台秤：1材质/2台面规格/3量程(分度值e=d)/4传感器*4/5接线盒*1/
            6仪表*1/7信号电缆5米/备注行
    台秤：  [防爆行]/1材质/2台面尺寸/3量程(e=d)/4传感器*1/5仪表*1/备注行

    lang=en 时输出 中文整块\\n英文整块(参照阿根廷双语范例)，英文材料/
    术语取自阿根廷语料；通表属性(量程/台面/材质)数值保持原样。
    """
    p_type = item["type"]
    row = (lookup_platform_row(item.get("platform"))
           if p_type == "platform" else lookup_bench_row(item.get("bench")))
    model = item.get("platform") or item.get("bench") or ""
    ex = bool(item.get("ex"))

    def _en_mtl(mtl):
        return _SCALE_MTL_TEXT_EN.get(mtl, mtl)

    def _en_sensor(text):
        return {"不锈钢": "stainless steel", "合金钢": "alloy steel"}.get(
            text, text)

    if row is None:
        # 通表无行(手输型号) -> 退化为简单型号行，保住可报价性
        zh = "型号：{}\n备注：".format(model) if with_model else "备注："
        if lang != "en":
            return zh
        en = "Model: {}\nRemarks:".format(model) if with_model else "Remarks:"
        return "{}\n{}".format(zh, en)

    # ---- 中文块(维持既有措辞) ----
    zh_lines = ["型号：{}".format(_scale_model_text(item))] if with_model else []

    if ex and p_type == "bench":
        # 台秤：防爆整机型号线(真实台秤单据第2行)；
        # 平台秤的防爆型号名移至名称列(B列)，描述不再重复
        zh_lines.append("防爆电子台秤")

    size = str(row.get("台面尺寸") or "").replace("x", "*").replace("X", "*")
    cap = row.get("kg量程") or ""
    zh_lines.append("1.材质：{}；".format(_SCALE_MTL_TEXT.get(
        (row.get("材质") or "").strip(), row.get("材质") or "")))
    zh_lines.append("2.台面{}：{}mm；".format(
        "规格" if p_type == "platform" else "尺寸", size))

    # 分度值e=d单一数值(平台秤不再单独列显示分度值d)
    if p_type == "platform":
        zh_lines.append("3.量程：{}kg，分度值e=d={}kg，精度等级C3；".format(
            cap, _num(row.get("kg分度值"))))
    else:
        zh_lines.append("3.量程：{}kg，分度值e=d={}kg，精度等级C3；".format(
            cap, _num(float(row.get("g分度值") or 0) / 1000)))

    brand = _BRAND_NAMES.get((row.get("品牌") or "").strip(), "")

    if p_type == "platform":
        sensors = (row.get("可用传感器") or "").split("|")
        sensor = sensors[0].strip() if sensors else ""
        zh_lines.append("4.称重传感器{}*4（材质：{}）；品牌：{}".format(
            sensor, sensor_material(sensor), brand))
        jbox = item.get("jbox") or ""
        jbrand = _BRAND_NAMES.get(
            (_JBOXES.get(_norm(jbox)) or {}).get("品牌", ""), "") if jbox else ""
        zh_lines.append("5.接线盒{}*1；品牌：{}".format(jbox, jbrand))
        ctrl_idx = 6
    else:
        zh_lines.append("4.{}传感器*1；品牌：{}".format("防爆" if ex else "", brand))
        ctrl_idx = 5

    controller = item.get("controller") or ""
    info = lookup_instrument(controller)
    desc = (info or {}).get("desc", "")
    install = item.get("controller_install") or ""
    cbrand = _BRAND_NAMES.get((info or {}).get("brand", ""), "")
    detail = install + desc
    # 描述通表文本已自带"；品牌：xxx"结尾时不再追加，避免品牌重复输出
    suffix = "" if "品牌：" in detail else (
        "；品牌：{}".format(cbrand) if cbrand else "")
    zh_lines.append("{}.{}仪表{}({})*1；{}{}".format(
        ctrl_idx, "防爆" if ex else "称重",
        instrument_display_prefix(controller), controller,
        detail, suffix,
    ))

    if p_type == "platform":
        zh_lines.append("7.信号电缆5米。")

    zh_lines.append("备注：")
    zh_text = "\n".join(zh_lines)

    if lang != "en":
        return zh_text

    # ---- 英文块(术语对齐阿根廷双语范例/通表 描述EN 语料) ----
    zh_model_text = _scale_model_text(item)
    en_model_text = zh_model_text.replace("（", "(").replace("）", ")")
    en_lines = []
    if with_model:
        en_lines.append("Model: {}".format(en_model_text))

    if ex and p_type == "bench":
        en_lines.append("Explosion-proof bench scale")

    mtl = (row.get("材质") or "").strip()
    en_lines.append("1. Material: {};".format(_en_mtl(mtl)))
    en_lines.append("2. Deck size: {}mm;".format(size))
    if p_type == "platform":
        en_lines.append(
            "3. Capacity: {}kg, certified interval e=d={}kg, accuracy class C3;".format(
                cap, _num(row.get("kg分度值"))))
    else:
        en_lines.append(
            "3. Capacity: {}kg, certified interval e=d={}kg, accuracy class C3;".format(
                cap, _num(float(row.get("g分度值") or 0) / 1000)))

    if p_type == "platform":
        sensors = (row.get("可用传感器") or "").split("|")
        sensor = sensors[0].strip() if sensors else ""
        en_lines.append("4. Load cells {}*4 (material: {}); Brand: {}".format(
            sensor, _en_sensor(sensor_material(sensor)), brand))
        en_lines.append("5. Junction box {}*1; Brand: {}".format(jbox, jbrand))
    else:
        en_lines.append("4. {}load cell*1; Brand: {}".format(
            "Explosion-proof " if ex else "", brand))

    desc_en = lookup_desc_en(controller)
    en_detail = desc_en
    if en_detail and cbrand and "brand:" not in en_detail.lower():
        en_detail = "{}; Brand: {}".format(en_detail, cbrand)
    ctrl_kind = "Explosion-proof " if ex else "Weighing "
    en_lines.append("{}. {}indicator {}({})*1; {}".format(
        ctrl_idx, ctrl_kind,
        instrument_display_prefix(controller), controller, en_detail,
    ))

    if p_type == "platform":
        en_lines.append("7. Signal cable 5 m.")

    en_lines.append("Remarks:")
    return "{}\n{}".format(zh_text, "\n".join(en_lines))


# ---------------------------------------------------------------------------
# 项目 -> 明细行模型  row: (名称/型号文本, 单位, 数量|None, 描述)
# ---------------------------------------------------------------------------

def _controller_row(controller_model, install=None, lang="zh"):
    """返回 (名称, 型号, 单位, 数量, 描述)；多通道仪表返回 (None, channels)
    由调用方汇总到底部。install 为支架安装前缀(立杆安装/壁挂安装)时插入描述开头。"""
    info = lookup_instrument(controller_model)
    desc_en = lookup_desc_en(controller_model)

    if info is None:
        return (_name_text(lang, "controller"), controller_model,
                _unit_text(lang, "pc"), 1, ""), 1

    name = _name_text(lang, "controller")

    if info["channels"] > 1:
        return None, info["channels"]

    zh_desc = install + info["desc"] if install else info["desc"]
    return (name, controller_model, _unit_text(lang, "pc"), 1,
            _desc_bi(lang, zh_desc, desc_en)), 1


def _guess_brand(item):
    """按已选型号判断器材品牌：AW模块(M开头)/AW平台秤/AW台秤/AW接线盒 -> AW，否则 FT"""
    for key in ("module", "platform", "bench"):
        value = item.get(key)
        model = value[0] if isinstance(value, tuple) else value

        if model:
            return "AW" if str(model).upper().startswith(("M", "AW")) else "FT"

    jbox = str(item.get("jbox") or "").upper()
    return "AW" if jbox.startswith("AW") else "FT"


def _append_controller(rows, tail_channels, item, qty, lang="zh"):
    """按规则把仪表行放入项目：非多通道 -> 行内数量1；
    多通道 -> 不入段，累计套数由底部统一行汇总"""
    controller = item.get("controller")
    line, channels = _controller_row(
        controller, item.get("controller_install"), lang)

    if line is not None:
        rows.append(line)
    elif channels > 1:
        tail_channels[controller] = tail_channels.get(controller, 0) + qty


# 电缆数量约定：模块业务标配10米(独立明细行)，仪表为多通道时20米；
# 平台秤/台秤标配5米，不另起一行，拼接进秤体描述(参照整机描述"…信号电缆5米；"写法)
_CABLE_QTY_MODULE = 10
_CABLE_QTY_MODULE_MULTI = 20
_CABLE_QTY_SCALE = 5

# engine仪表行 com_2 代码(parser rules protocols canonical) -> 人类语言全称，
# 分项段header的"仪表，xxx通信"用
_COM2_FULL = {
    "ANALOG": "4-20mA模拟量输出",
    "TCP": "Modbus-TCP以太网",
    "UDP": "Modbus-UDP以太网",
    "DP": "Profibus-DP",
    "PN": "Profinet",
    "PROFINET": "Profinet",
    "IP": "EtherNet-IP",
    "CAT": "EtherCAT",
}

# com_2 代码 -> 英文全称(英文报价分项段header用)
_COM2_FULL_EN = {
    "ANALOG": "4-20 mA analog output",
    "TCP": "Modbus-TCP Ethernet",
    "UDP": "Modbus-UDP Ethernet",
    "DP": "Profibus-DP",
    "PN": "Profinet",
    "PROFINET": "Profinet",
    "IP": "EtherNet-IP",
    "CAT": "EtherCAT",
}


def com2_full(code, lang="zh"):
    """com_2 代码(如 'TCP|UDP'/'ANALOG'/0) -> 全称；无扩展通讯返回空串"""
    token = str(code or "").split("|")[0].strip().upper()
    table = _COM2_FULL_EN if lang == "en" else _COM2_FULL
    return table.get(token, "")


def _num(value):
    """数值 -> 去尾零文本(3400->'3400', 0.5->'0.5')；非法/None返回''"""
    try:
        text = ("%.6f" % float(value)).rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        return ""
    return text


def _t(kg):
    """kg数值 -> 吨文本(3400->'3.4')；非法/None返回''"""
    try:
        return _num(float(kg) / 1000)
    except (TypeError, ValueError):
        return ""


def _instrument_segments(h, lang="zh"):
    """仪表段片段：[com_2全称(通信), 供电方式]，无扩展通讯时仅供电"""
    comm = com2_full(h.get("com2"), lang)
    comm_text = (comm if lang == "en" else "{}通信".format(comm)) if comm else ""
    return [s for s in (comm_text, str(h.get("power") or "")) if s]


def build_multichannel_header(items, model, lang="zh"):
    """多通道仪表分项段标题：拼接逻辑与模块header的仪表段一致。
    多通道仪表，{com_2全称}通信，{供电方式}。{防爆环境。}
    com_2/供电/防爆取自选用该仪表的项目(同一型号的仪表属性一致，取首个匹配)。
    lang=en 时输出 中文\\n英文 双语。"""
    header, ex = {}, False

    for item in items or []:
        if item.get("controller") == model:
            header = item.get("header") or {}
            ex = bool(item.get("ex"))
            break

    zh = "多通道仪表，{}。".format(
        "，".join(_instrument_segments(header, "zh")))
    zh += _txt("zh", "ex_note") if ex else ""
    if lang != "en":
        return zh
    en = "Multi-channel controller, {}.".format(
        ", ".join(_instrument_segments(header, "en")))
    en += _txt("en", "ex_note") if ex else ""
    return "{}\n{}".format(zh, en)


def _header_instrument(h, lang="zh"):
    """仪表段：com_2全称通信 + 供电方式"""
    sep = ", " if lang == "en" else "，"
    text = "Controller, {}." if lang == "en" else "仪表，{}。"
    return text.format(sep.join(_instrument_segments(h, lang)))


def build_section_header(item, lang="zh"):
    """多项目模板分项段header文本(按业务pipeline分别拼接，Writer统一写入)：

    模块   ：设备名称： ，位号： ，N支点，共X套。设备自重：.. t，料重：.. t，
             总重.. t，标定量程.. t，显示分度值.. kg。仪表，..。防爆环境。
             (EN) Equipment name / S/N / supports / sets / tare-weight /
             rated capacity / display interval / Ex-proof environment，
             与英文模板占位文案同词表。
    平台秤 ：电子平台秤，共X台。量程..kg，分度值..kg。仪表，..。
    台秤   ：电子台秤，共X台。量程..kg，显示分度值..g。仪表，..。

    数据来源：W1/W2/支点/套数由 Checklist 从 Desk 采集带入；模块的标定量程/
    显示分度值取 engine metrology；秤业务的量程/分度值按用户确认型号回查通表
    (与整机描述同源)；仪表段 com_2全称/供电来自 engine 仪表行。
    """
    p_type = item["type"]
    h = item.get("header") or {}
    qty = item.get("qty")
    instrument = _header_instrument(h, lang)

    if p_type == "module":
        support = item.get("support")
        w1, w2 = h.get("W1"), h.get("W2")

        if isinstance(w1, (int, float, str)) and isinstance(
                w2, (int, float, str)):
            total = _t(float(w1) + float(w2))
        else:
            total = ""

        zh_text = "设备名称： ，位号： ，{}，共{}。设备自重：{} t，料重：{} t，总重{} t，标定量程{} t，显示分度值{} kg。{}".format(
            "{}支点".format(support) if support else " 支点",
            "{}{}".format(qty, "套") if qty is not None else " 套",
            _t(w2), _t(w1), total, _t(h.get("range")), _num(h.get("e")),
            _header_instrument(h, "zh"),
        )
        zh_text += _txt("zh", "ex_note") if item.get("ex") else ""

        if lang != "en":
            return zh_text

        support_text = ("{} supports".format(support)
                        if support else "  supports")
        qty_text = ("{} sets".format(qty)
                    if qty is not None else "  sets")
        en_text = ("Equipment name:  , S/N:  , {}, {}. "
                   "Equipment tare-weight: {} t, material weight: {} t, "
                   "total weight {} t, rated capacity: {} t, "
                   "display interval: {} kg. {}").format(
            support_text, qty_text, _t(w2), _t(w1), total,
            _t(h.get("range")), _num(h.get("e")),
            _header_instrument(h, "en"),
        )
        if item.get("ex"):
            en_text += " " + _txt("en", "ex_note")
        return "{}\n{}".format(zh_text, en_text)

    if p_type == "platform":
        row = lookup_platform_row(item.get("platform")) or {}
        zh_name = _txt("zh", "platform_ex" if item.get("ex") else "platform")
        en_name = _txt("en", "platform_ex" if item.get("ex") else "platform")

        zh_text = "{}，共{}{}。量程{}kg，分度值{}kg。{}".format(
            zh_name, qty if qty is not None else "", _txt("zh", "scale_unit"),
            row.get("kg量程") or "", _num(row.get("kg分度值")),
            _header_instrument(h, "zh"),
        )
        if lang != "en":
            return zh_text

        en_text = "{}, {} units in total. Capacity: {} kg, scale interval: {} kg. {}".format(
            en_name, qty if qty is not None else "",
            row.get("kg量程") or "", _num(row.get("kg分度值")),
            _header_instrument(h, "en"),
        )
        return "{}\n{}".format(zh_text, en_text)

    row = lookup_bench_row(item.get("bench")) or {}
    zh_text = "电子台秤，共{}{}。量程{}kg，显示分度值{}g。{}".format(
        qty if qty is not None else "", _txt("zh", "scale_unit"),
        row.get("kg量程") or "", _num(row.get("g分度值")),
        _header_instrument(h, "zh"),
    )
    if lang != "en":
        return zh_text

    en_text = "Electronic bench scale, {} units in total. Capacity: {} kg, display interval: {} g. {}".format(
        qty if qty is not None else "",
        row.get("kg量程") or "", _num(row.get("g分度值")),
        _header_instrument(h, "en"),
    )
    return "{}\n{}".format(zh_text, en_text)


def _with_cable_note(desc):
    """在秤体描述尾部(品牌后缀之前)拼接 信号电缆Nm"""
    note = "，信号电缆{}米".format(_CABLE_QTY_SCALE)

    if not desc:
        return note.lstrip("，")

    idx = desc.find("；品牌：")

    if idx > 0:
        return desc[:idx] + note + desc[idx:]

    return desc + note


# ---------------------------------------------------------------------------
# 过渡板/通讯电缆选型(模块业务)
# ---------------------------------------------------------------------------

# 传感器容量/单位，如 SLB-227kg / SB14-4536kg / AWB-5t
_PLATE_SENSOR_RE = re.compile(r"(\d+(?:\.\d+)?)(kg|t)\b")

# AW 模块材质段：M20-C- / M30L-S- / M50A-C- -> C(碳钢)/S(不锈钢)
_PLATE_AW_MATERIAL_RE = re.compile(r"\bM\d+A?L?-(C|S)-")


def _module_plate_model(module_spec):
    """模块完整规格型号 -> 过渡板型号("CS 3t"式)；解析失败返回""(调用方回退通用行)。

    材质跟随模块材质(FT: CS/SS token；AW: Mxx-C-/Mxx-S- 段)；
    档位优先按模块系列：55-30全系 22.5t(FT)、M50全系 20t(AW，含M50A)，
    其余按传感器容量——FT ≤3t→3t、4536kg→4.5t、更大(5099kg)→5t；
    AW ≤3t→2t、5t→5t。
    """
    spec = module_spec or ""

    material = next((t for t in spec.split() if t in ("CS", "SS")), None)
    m = _PLATE_AW_MATERIAL_RE.search(spec)

    if material is None and m:
        material = "CS" if m.group(1) == "C" else "SS"

    if material is None:
        return ""

    if spec.startswith("55-30"):
        rating = "22.5t"
    elif spec.startswith("M50"):
        rating = "20t"
    else:
        hit = _PLATE_SENSOR_RE.search(spec)

        if hit is None:
            return ""

        cap = float(hit.group(1)) * (1000 if hit.group(2) == "t" else 1)

        if spec.startswith("M"):
            rating = "5t" if cap > 3000 else "2t"
        elif cap <= 3000:
            rating = "3t"
        else:
            rating = "4.5t" if cap <= 4536 else "5t"

    return "{} {}".format(material, rating)


def build_item_rows(item, scale_model_in_name=False, lang="zh"):
    """单个项目 -> (明细行列表, 底部追加的多通道仪表 {model: 累计套数})。
    行: (名称, 型号, 单位, 数量, 描述)；lang=en 时名称为"英文\\n中文"、
    单位为"英文\\n中文"、描述为"中文在上英文下"(描述EN 缺失回退中文)。
    scale_model_in_name=True(COPA B多项目模板 B=名称/型号)：型号由调用方并入
    名称列，描述不含型号行；否则型号由调用方并入描述首行(COPA A多项目/内容块)。"""
    rows, tail_channels = [], {}
    p_type = item["type"]
    qty = item.get("qty") or 1

    if p_type == "module":
        # 通讯电缆：数字化方案统一 AWG-14/5D(描述优先取描述通表，未录入时
        # 兜底固定文案)；模拟方案品牌随模块品牌；仪表为多通道时电缆加长配20米
        if item.get("digital"):
            cable = "AWG-14/5D"
            cable_desc = lookup_desc(cable) or "4芯屏蔽电缆"
        else:
            cable = "AWP-06/3" if _guess_brand(item) == "AW" else "AWG-26/3"
            cable_desc = lookup_desc(cable)
        info = lookup_instrument(item.get("controller"))
        cable_qty = (_CABLE_QTY_MODULE_MULTI
                     if info is not None and info["channels"] > 1
                     else _CABLE_QTY_MODULE)
        support = item.get("support") or 4
        model, desc = item.get("module") or (None, "")
        rows.append((_name_text(lang, "module"), model or "",
                     _unit_text(lang, "pc"), support,
                     _desc_bi(lang, desc, lookup_desc_en(model))))
        jbox = item.get("jbox")
        rows.append((_name_text(lang, "jbox"), jbox or "",
                     _unit_text(lang, "pc"), 1,
                     _desc_bi(lang, lookup_accessory(jbox),
                              lookup_desc_en(jbox))))
        # 单/多通道仪表按模板行序放在接线盒之后
        _append_controller(rows, tail_channels, item, qty, lang)
        rows.append((_name_text(lang, "cable"), cable,
                     _unit_text(lang, "m"), cable_qty,
                     _desc_bi(lang, cable_desc, lookup_desc_en(cable))))
        # 过渡板按附件选择输出(上/下型号描述相同，行首标明方向)；
        # 型号随模块材质与传感器系列/容量解析
        accessories = item.get("accessories") or []
        plate = _module_plate_model(model or "")

        for key, option in (("plate_up", "上过渡板"),
                            ("plate_down", "下过渡板")):
            if option not in accessories:
                continue

            if plate:
                rows.append((_name_text(lang, key), plate,
                             _unit_text(lang, "plate"), support,
                             _desc_bi(lang, lookup_desc(plate),
                                      lookup_desc_en(plate))))
            else:
                rows.append((_name_text(lang, key), "",
                             _unit_text(lang, "plate"), support,
                             lookup_accessory("过渡板")))
    elif p_type == "platform":
        # 整机名称：防爆平台秤为 防爆电子平台秤(B列名称)；整机描述自带
        # 型号行(with_model=True)时型号不再单独落型号列
        base = "platform_ex" if item.get("ex") else "platform"
        with_model = not scale_model_in_name
        model = "" if with_model else _scale_model_text(item)
        rows.append((_name_text(lang, base), model,
                     _unit_text(lang, "scale_unit"), qty,
                     build_scale_desc(item, with_model=with_model, lang=lang)))
    elif p_type == "bench":
        base = "bench"
        with_model = not scale_model_in_name
        model = "" if with_model else _scale_model_text(item)
        rows.append((_name_text(lang, base), model,
                     _unit_text(lang, "scale_unit"), qty,
                     build_scale_desc(item, with_model=with_model, lang=lang)))

    # 大屏幕：附件选型确认的具体型号(Checklist display 字段)落行为只数1
    display = item.get("display")

    if display:
        name = (_name_text(lang, "display_ex") if display.startswith("AWRD")
                else _name_text(lang, "display"))
        rows.append((name, display, _unit_text(lang, "pc"), 1,
                     _desc_bi(lang, lookup_desc(display),
                              lookup_desc_en(display))))

    return rows, tail_channels


# ---------------------------------------------------------------------------
# 模板填充
# ---------------------------------------------------------------------------

def _style_of(cell):
    return (copy(cell.font), copy(cell.fill), copy(cell.border),
            copy(cell.alignment), cell.number_format)


def _apply_style(cell, style):
    font, fill, border, alignment, number_format = style
    cell.font, cell.fill, cell.border = font, fill, border
    cell.alignment, cell.number_format = alignment, number_format


def _capture_row(ws, row, col_start, col_end):
    """快照一行：每格 (值, 样式) + 行高，供后续重写/复制"""
    return {
        "height": ws.row_dimensions[row].height,
        "cells": {
            col: (ws.cell(row=row, column=col).value, _style_of(ws.cell(row=row, column=col)))
            for col in range(col_start, col_end + 1)
        },
    }


def _set_cell(ws, row, col, value):
    """写单元格值。ws.cell() 的静态类型含 MergedCell(其 value 属性只读)，
    直接赋值会报 "Cannot assign to attribute value"；统一经此助手写入。"""
    cast(Any, ws.cell(row=row, column=col)).value = value


def _write_row(ws, row, proto, values, col_start=1):
    """按原型行样式写值；values: {列号: 值}，未给出的列沿用原型值"""
    ws.row_dimensions[row].height = proto["height"]

    for col, (default, style) in proto["cells"].items():
        cell = cast(Any, ws.cell(row=row, column=col))
        cell.value = values.get(col, default)
        _apply_style(cell, style)


def _stamp_date(ws):
    """把抬头区的 DATE：/报价日期：/DATE:/Quotation date: 改写为输出当天日期
    (YYYY-MM-DD)。COPA B模板在 A10、COPA A在 A9(中英文模板前缀不同)；
    找不到时静默跳过(模板改动无碍)。"""
    today = datetime.now().strftime("%Y-%m-%d")

    for row in range(1, 12):
        value = ws.cell(row=row, column=1).value

        if not isinstance(value, str):
            continue

        for prefix in ("DATE：", "报价日期：", "DATE:", "Quotation date:"):
            if value.startswith(prefix):
                _set_cell(ws, row, 1, prefix + today)
                return


def _stamp_company_contact(ws, items):
    """把公司名称/联系人写入抬头询价区(值由 Checklist 随项目dict下发)。
    COPA B模板为 A7 TO：/A8 ATTN：，COPA A为 询价方:/To: 与 联系人：/Attn:
    (中英文前缀不同)；栏位找不到或字段为空时静默跳过(模板改动无碍)。"""
    first = items[0] if items else {}
    pairs = (
        (("TO：", "TO:", "To:", "询价方:"),
         str(first.get("company") or "").strip()),
        (("ATTN：", "ATTN:", "Attn:", "联系人："),
         str(first.get("contact") or "").strip()),
    )

    if not any(data for _, data in pairs):
        return

    for row in range(1, 12):
        value = ws.cell(row=row, column=1).value

        if not isinstance(value, str):
            continue

        for prefixes, data in pairs:
            if data and value.startswith(prefixes):
                _set_cell(ws, row, 1, value + data)


def _write_single(items, out_path, service=False, lang="zh"):
    """COPA B单项目报价模板：明细区 13-19 行(7行)，其后为服务费行(仅中文模板含)
    与价税合计行。现行业务最大 6 行明细(+1 多通道底部行)，无需插行，避免
    openpyxl 行增删不迁移公式/合并单元格的问题；未用行清空。
    服务费/价税合计行按内容锚点定位(模板备注区/汇总区可自由增删)。"""
    tpl = os.path.join(TEMPLATE_DIR, _tpl_name("COPA B单项目报价模板.xlsx", lang))
    wb = load_workbook(tpl)
    ws = wb.worksheets[0]
    _stamp_date(ws)
    _stamp_company_contact(ws, items)
    rows, tail = build_item_rows(items[0], lang=lang)
    # 多通道仪表 -> 从项目删去，底部统一追加：数量 = ceil(spinbox ÷ 通道数)
    tail_rows = []

    for model, sets in tail.items():
        info = lookup_instrument(model)
        channels = info["channels"] if info else 4
        tail_rows.append((
            _name_text(lang, "controller"), model,
            _unit_text(lang, "pc"),
            math.ceil(sets / channels),
            _desc_bi(lang, info["desc"] if info else "",
                     lookup_desc_en(model)),
        ))

    all_rows = rows + tail_rows
    proto = _capture_row(ws, 13, 1, 8)

    if len(all_rows) > 7:
        raise ValueError("单项目明细超过7行({}行)，请改用多项目模板或精简项目".format(len(all_rows)))

    for i in range(7):
        r = 13 + i

        if i < len(all_rows):
            name, model, unit, qty, desc = all_rows[i]
            spec = "\n".join(x for x in (model, desc) if x)
            _write_row(ws, r, proto, {
                1: "=ROW()-12",        # 序号
                2: name,               # 名称
                3: spec,               # 型号规格、描述
                4: unit, 5: qty, 7: "=E{0}*F{0}".format(r),
            })
        else:
            _write_row(ws, r, proto, {1: None, 2: None, 3: None, 4: None, 5: None, 6: None, 7: None})

    # 服务费/价税合计行按内容定位：zh 模板含服务费行(现场服务开=保留占位，
    # 关=整行清空并收缩合计公式)；en 模板无服务费行，开关不改变输出结构
    r_serve = next((r for r in range(13, ws.max_row + 1)
                    if str(ws.cell(row=r, column=2).value or "").startswith("服务费")),
                   None)
    r_total = next((r for r in range(13, ws.max_row + 1)
                    if str(ws.cell(row=r, column=1).value or "").startswith(
                        ("价税合计", "Total"))),
                   None)

    if r_serve is not None and not service:
        _write_row(ws, r_serve, proto,
                   {1: None, 2: None, 3: None, 4: None, 5: None, 6: None, 7: None})

        if r_total is not None:
            _set_cell(ws, r_total, 7, "=SUM(G13:G19)")

    if service and r_serve is None:
        print("[Writer] {}：模板无服务费行，现场服务开关未生效".format(
            os.path.basename(out_path)))

    wb.properties.creator = "Z.ai"
    wb.save(out_path)


def _write_multi(items, out_path, service=False, lang="zh"):
    """COPA B多项目报价模板：每个 quote_id 一个分项段(标题+明细+小计+套数合计)，
    段后可接多通道仪表汇总行，然后 设备合计/服务费用/总价/备注。
    区域 12 行起整体重建：openpyxl 不迁移公式与合并单元格，故全部显式重写。"""
    tpl = os.path.join(TEMPLATE_DIR, _tpl_name("COPA B多项目报价模板.xlsx", lang))
    wb = load_workbook(tpl)
    ws = wb.worksheets[0]
    _stamp_date(ws)
    _stamp_company_contact(ws, items)

    # 1) 快照原模板 12 行起(值/样式/行高)与区内合并单元格；备注区行数随模板
    protos = {r: _capture_row(ws, r, 1, 8) for r in range(12, ws.max_row + 1)}
    merges = [(m.min_row, m.min_col, m.max_col) for m in ws.merged_cells.ranges if m.min_row >= 12]
    merge_by_row = {}

    for row, c1, c2 in merges:
        merge_by_row.setdefault(row, []).append((c1, c2))

    for m in [m for m in list(ws.merged_cells.ranges) if m.min_row >= 12]:
        ws.unmerge_cells(str(m))

    for r in range(12, ws.max_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).value = None

    # 2) 原型：12段标题 13明细 19小计 20套数合计；汇总/备注区按内容锚点定位
    #    (A列'A'=设备合计，其后可有服务费行，再总价行，备注头+备注至表尾)
    P_TITLE, P_ITEM, P_SUB, P_SETS = protos[12], protos[13], protos[19], protos[20]
    r_sum_src = next(r for r in protos
                     if protos[r]["cells"].get(1, (None,))[0] == "A")
    r_serve_src = (r_sum_src + 1
                   if str(protos[r_sum_src + 1]["cells"].get(2, (None,))[0]
                          or "").startswith("服务费")
                   else None)
    r_total_src = (r_serve_src or r_sum_src) + 1
    P_SUM, P_TOTAL = protos[r_sum_src], protos[r_total_src]

    # 3) 逐项目写分项段
    row = 12
    section_spans = []      # (首行, 末行)
    tail_total = {}         # 多通道仪表 model -> 累计套数

    for idx, item in enumerate(items):
        # COPA B多项目模板 B=名称/型号：秤业务型号并入名称列(描述去型号行)
        rows, tail = build_item_rows(item, scale_model_in_name=True,
                                     lang=lang)

        for model, sets in tail.items():
            tail_total[model] = tail_total.get(model, 0) + sets

        unit_word = (_txt(lang, "set") if item["type"] == "module"
                     else _txt(lang, "scale_unit"))
        first_item = row + 1
        _write_row(ws, row, P_TITLE, {
            1: _unit_names(lang)[idx] if idx < 10 else str(idx + 1),
            2: build_section_header(item, lang),
        })
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

        for i, (name, model, unit, qty, desc) in enumerate(rows):
            r = first_item + i
            # COPA B多项目模板 B=名称/型号：型号行并入名称格
            combined = "\n".join(x for x in (name, model) if x)
            _write_row(ws, r, P_ITEM, {
                1: i + 1, 2: combined, 3: unit, 4: qty, 7: desc,
                6: "=D{0}*E{0}".format(r),
            })

        last_item = first_item + len(rows) - 1
        r_sub, r_sets = last_item + 1, last_item + 2
        _write_row(ws, r_sub, P_SUB, {1: len(rows) + 1, 6: "=SUM(F{}:F{})".format(first_item, last_item)})
        _write_row(ws, r_sets, P_SETS, {
            1: len(rows) + 2, 3: unit_word, 4: item.get("qty") or 1,
            6: "=F{}*D{}".format(r_sub, r_sets),
        })
        section_spans.append((first_item, r_sets))
        row = r_sets + 1

    # 4) 多通道仪表 -> 项目底部统一追加独立分项段：
    #    段标题 + 仪表明细行 + 小计。不写套数合计(仪表不影响套数统计)，
    #    小计金额另行并入总价
    tail_sub_rows = []

    for model, sets in tail_total.items():
        info = lookup_instrument(model)
        channels = info["channels"] if info else 4
        qty = math.ceil(sets / channels)
        names = _unit_names(lang)
        numeral = names[len(items)] if len(items) < 10 else str(len(items) + 1)
        _write_row(ws, row, P_TITLE, {
            1: numeral, 2: build_multichannel_header(items, model, lang),
        })
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        row += 1
        r_item = row
        _write_row(ws, row, P_ITEM, {
            1: 1,
            2: "\n".join(x for x in (
                _name_text(lang, "controller"), model) if x),
            3: _unit_text(lang, "pc"),
            4: qty,
            7: _desc_bi(lang, info["desc"] if info else "",
                        lookup_desc_en(model)),
            6: "=D{0}*E{0}".format(row),
        })
        row += 1
        _write_row(ws, row, P_SUB, {
            1: 2, 6: "=SUM(F{}:F{})".format(r_item, r_item),
        })
        tail_sub_rows.append(row)
        row += 1

    # 5) 汇总行(设备合计/[服务费]/总价)：现场服务开关决定服务费行——
    #    zh 模板含服务费行(开=数量1单位次并计入总价；关=整行跳过)；
    #    en 模板无服务费行，开关不改变输出结构(CLI 提示)。
    #    多通道仪表小计并入总价，不影响设备合计的套数统计
    span_top = section_spans[0][0] - 1
    span_bottom = section_spans[-1][1]
    tail_extra = "".join("+F{}".format(r) for r in tail_sub_rows)
    r_sum = row
    sets_total = _txt(lang, "sets_total")
    _write_row(ws, r_sum, P_SUM, {
        4: '=SUMIF(B{}:B{},"{}",D{}:D{})'.format(
            span_top, span_bottom, sets_total, span_top, span_bottom),
        6: '=SUMIF(B{}:B{},"{}",F{}:F{})'.format(
            span_top, span_bottom, sets_total, span_top, span_bottom),
    })

    if service and r_serve_src is not None:
        # 开 -> 服务费行保留模板占位(数量1/次/描述来自模板，writer不写业务值)，
        # 仅补总价结构公式让填价即计总；多通道仪表小计并入总价
        _write_row(ws, r_sum + 1, protos[r_serve_src],
                   {6: "=D{0}*E{0}".format(r_sum + 1)})
        _write_row(ws, r_sum + 2, P_TOTAL, {
            4: "=D{}".format(r_sum),
            6: "=F{}+F{}{}".format(r_sum, r_sum + 1, tail_extra),
        })
        row = r_sum + 3
    else:
        if service:
            print("[Writer] {}：模板无服务费行，现场服务开关未生效".format(
                os.path.basename(out_path)))

        # 关 -> 服务费整行跳过，总价=设备合计+多通道仪表小计
        _write_row(ws, r_sum + 1, P_TOTAL, {
            4: "=D{}".format(r_sum), 6: "=F{}{}".format(r_sum, tail_extra),
        })
        row = r_sum + 2

    # 6) 备注区(备注头 + 备注，行数随模板)平移重写
    note_start = row
    _write_row(ws, row, protos[r_total_src + 1], {})
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
    row += 1

    for r in range(r_total_src + 2, ws.max_row + 1):
        a_val = protos[r]["cells"].get(1, (None, None))[0]
        b_val = protos[r]["cells"].get(2, (None, None))[0]

        if a_val is None and b_val is None:
            continue

        _write_row(ws, row, protos[r], {1: a_val, 2: b_val})

        for c1, c2 in merge_by_row.get(r, []):
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)

        row += 1

    wb.properties.creator = "Z.ai"
    wb.save(out_path)


# ---------------------------------------------------------------------------
# COPA A模板填充(列布局：A序号 B名称 C:D型号规格描述 E单位 F数量 G单价 H总价 I备注)
# 与COPA B的差异：明细区行数不足需整体重建；价税合计行带大写金额公式，
# 其 CONCATENATE 长公式引用总价单元格，重建后必须把旧引用替换为新行号。
# ---------------------------------------------------------------------------

def _rebuild_area(ws, first_row):
    """快照 first_row 起全部行(值/样式/行高)与合并单元格，然后解合并并清空值。
    返回 (protos, merge_by_row)。openpyxl 增删行不迁移公式/合并，故走重建路线。"""
    last_row = ws.max_row
    protos = {r: _capture_row(ws, r, 1, 9) for r in range(first_row, last_row + 1)}
    merge_by_row = {}

    for m in list(ws.merged_cells.ranges):
        if m.min_row >= first_row:
            merge_by_row.setdefault(m.min_row, []).append((m.min_col, m.max_col))
            ws.unmerge_cells(str(m))

    for r in range(first_row, last_row + 1):
        for c in range(1, 10):
            ws.cell(row=r, column=c).value = None

    return protos, merge_by_row


def _tail_rows(tail_total, lang="zh"):
    """多通道仪表汇总行：数量 = ceil(该仪表覆盖项目套数合计 ÷ 通道数)"""
    result = []

    for model, sets in tail_total.items():
        info = lookup_instrument(model)
        channels = info["channels"] if info else 4
        result.append((
            _name_text(lang, "controller"), model,
            _unit_text(lang, "pc"),
            math.ceil(sets / channels),
            _desc_bi(lang, info["desc"] if info else "",
                     lookup_desc_en(model)),
        ))

    return result


def _write_aw_single(items, out_path, service=False, lang="zh"):
    """COPA A单项目报价模板：明细区 12 行起动态重建；
    服务费占位行(仅中文模板含)按内容锚点定位，开=保留占位仅重映射公式，
    关=整行删除；价税合计行大写公式引用原合计行号，需替换为实际行号。"""
    tpl = os.path.join(TEMPLATE_DIR, _tpl_name("COPA A单项目报价模板.xlsx", lang))
    wb = load_workbook(tpl)
    ws = wb.worksheets[0]
    _stamp_date(ws)
    _stamp_company_contact(ws, items)
    rows, tail = build_item_rows(items[0], lang=lang)
    all_rows = rows + _tail_rows(tail, lang)

    protos, merge_by_row = _rebuild_area(ws, 12)
    P_ITEM = protos[12]
    r_serve_src = next((r for r in protos
                        if str(protos[r]["cells"].get(2, (None,))[0] or ""
                               ).startswith("服务费")),
                       None)
    r_total_src = next(r for r in protos
                       if str(protos[r]["cells"].get(1, (None,))[0] or ""
                              ).startswith(("价税合计", "Grand Total")))
    P_TOTAL = protos[r_total_src]

    row = 12

    for name, model, unit, qty, desc in all_rows:
        spec = "\n".join(x for x in (model, desc) if x)
        _write_row(ws, row, P_ITEM, {
            1: "=ROW()-11", 2: name, 3: spec,
            5: unit, 6: qty, 8: "=F{0}*G{0}".format(row),
        })
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        row += 1

    # 现场服务开 -> 服务费行保留模板占位(数量1/次/描述均取自模板占位行)，
    # writer 只重映射序号/总价公式；关 -> 整行删除(不写)，合计公式随之下移不含它；
    # en 模板无服务费行，开关不改变输出结构
    if service and r_serve_src is not None:
        _write_row(ws, row, protos[r_serve_src], {
            1: "=ROW()-11", 8: "=F{0}*G{0}".format(row),
        })

        for c1, c2 in merge_by_row.get(r_serve_src, []):
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)

        row += 1
    elif service:
        print("[Writer] {}：模板无服务费行，现场服务开关未生效".format(
            os.path.basename(out_path)))

    last_item = row - 1
    r_total = row
    cap_formula = str(P_TOTAL["cells"][3][0] or "").replace(
        "H{}".format(r_total_src), "H{}".format(r_total))
    _write_row(ws, r_total, P_TOTAL, {
        3: cap_formula, 8: "=SUM(H12:H{})".format(last_item),
    })
    ws.merge_cells(start_row=r_total, start_column=3, end_row=r_total, end_column=6)
    ws.merge_cells(start_row=r_total, start_column=8, end_row=r_total, end_column=9)
    row = r_total + 1

    # 报价说明/TERMS 头 + 备注(行数随模板)
    for r in range(r_total_src + 1, max(protos) + 1):
        a_val = protos[r]["cells"].get(1, (None, None))[0]
        b_val = protos[r]["cells"].get(2, (None, None))[0]

        if a_val is None and b_val is None:
            continue

        _write_row(ws, row, protos[r], {1: a_val, 2: b_val})

        for c1, c2 in merge_by_row.get(r, []):
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)

        row += 1

    wb.properties.creator = "Z.ai"
    wb.save(out_path)


def _write_aw_multi(items, out_path, service=False, lang="zh"):
    """COPA A多项目报价模板：每个 quote_id 一个分项段(标题B:I合并、明细C:D合并、
    小计、套数合计)，可接多通道仪表汇总行，后接 设备合计/服务费/价税合计(大写)/报价说明。"""
    tpl = os.path.join(TEMPLATE_DIR, _tpl_name("COPA A多项目报价模板.xlsx", lang))
    wb = load_workbook(tpl)
    ws = wb.worksheets[0]
    _stamp_date(ws)
    _stamp_company_contact(ws, items)
    protos, merge_by_row = _rebuild_area(ws, 12)

    # 原型：12段标题 13明细 19小计 20套数合计；汇总/备注区按内容锚点定位
    #    (A列'A'=设备合计，其后可有服务费行，再价税合计行，说明头+备注至表尾)
    P_TITLE, P_ITEM, P_SUB, P_SETS = protos[12], protos[13], protos[19], protos[20]
    r_sum_src = next(r for r in protos
                     if protos[r]["cells"].get(1, (None,))[0] == "A")
    r_serve_src = (r_sum_src + 1
                   if str(protos[r_sum_src + 1]["cells"].get(2, (None,))[0]
                          or "").startswith("服务费")
                   else None)
    r_total_src = (r_serve_src or r_sum_src) + 1
    P_SUM, P_TOTAL = protos[r_sum_src], protos[r_total_src]

    row = 12
    section_spans = []
    tail_total = {}

    for idx, item in enumerate(items):
        rows, tail = build_item_rows(item, lang=lang)

        for model, sets in tail.items():
            tail_total[model] = tail_total.get(model, 0) + sets

        unit_word = (_txt(lang, "set") if item["type"] == "module"
                     else _txt(lang, "scale_unit"))
        first_item = row + 1
        numeral = _unit_names(lang)[idx] if idx < 10 else str(idx + 1)
        _write_row(ws, row, P_TITLE, {
            1: numeral + ("." if lang == "en" else "、"),
            2: build_section_header(item, lang),
        })
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)

        for i, (name, model, unit, qty, desc) in enumerate(rows):
            r = first_item + i
            spec = "\n".join(x for x in (model, desc) if x)
            _write_row(ws, r, P_ITEM, {
                1: i + 1, 2: name, 3: spec, 5: unit, 6: qty,
                8: "=F{0}*G{0}".format(r),
            })
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)

        last_item = first_item + len(rows) - 1
        r_sub, r_sets = last_item + 1, last_item + 2
        _write_row(ws, r_sub, P_SUB, {
            1: len(rows) + 1, 5: _txt(lang, "set"),
            8: "=SUM(H{}:H{})".format(first_item, last_item),
        })
        _write_row(ws, r_sets, P_SETS, {
            1: len(rows) + 2, 5: unit_word, 6: item.get("qty") or 1,
            8: "=F{}*H{}".format(r_sets, r_sub),
        })
        section_spans.append((first_item, r_sets))
        row = r_sets + 1

    # 多通道仪表 -> 项目底部统一追加行(序号沿用段式中文数字)
    # 多通道仪表 -> 项目底部统一追加独立分项段(与项目段结构统一)：
    # 段标题 + 仪表明细行 + 小计。不写套数合计(仪表不影响套数统计)，
    # 小计金额另行并入价税合计
    tail_sub_rows = []

    for model, sets in tail_total.items():
        info = lookup_instrument(model)
        channels = info["channels"] if info else 4
        qty = math.ceil(sets / channels)
        names = _unit_names(lang)
        numeral = names[len(items)] if len(items) < 10 else str(len(items) + 1)
        _write_row(ws, row, P_TITLE, {
            1: numeral + ("." if lang == "en" else "、"),
            2: build_multichannel_header(items, model, lang),
        })
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        row += 1
        r_item = row
        _write_row(ws, row, P_ITEM, {
            1: 1, 2: _name_text(lang, "controller"),
            3: "\n".join(x for x in (
                model, info["desc"] if info else "",
                lookup_desc_en(model)) if x),
            5: _unit_text(lang, "pc"), 6: qty,
            8: "=F{0}*G{0}".format(row),
        })
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        row += 1
        _write_row(ws, row, P_SUB, {
            1: 2, 5: _txt(lang, "set"), 8: "=SUM(H{}:H{})".format(r_item, r_item),
        })
        tail_sub_rows.append(row)
        row += 1

    # 设备合计 / [服务费] / 价税合计(大写公式引用需替换)。
    # zh 模板含服务费行(开=数量1/单位次并计入总价；关=跳过)；
    # en 模板无服务费行，开关不改变输出结构(CLI 提示)。
    span_top = section_spans[0][0] - 1
    span_bottom = section_spans[-1][1]
    tail_extra = "".join("+H{}".format(r) for r in tail_sub_rows)
    r_sum = row
    sets_total = _txt(lang, "sets_total")
    _write_row(ws, r_sum, P_SUM, {
        6: '=SUMIF(B{}:B{},"{}",F{}:F{})'.format(
            span_top, span_bottom, sets_total, span_top, span_bottom),
        8: '=SUMIF(B{}:B{},"{}",H{}:H{})'.format(
            span_top, span_bottom, sets_total, span_top, span_bottom),
    })

    if service and r_serve_src is not None:
        _write_row(ws, r_sum + 1, protos[r_serve_src],
                   {8: "=F{0}*G{0}".format(r_sum + 1)})
        r_total = r_sum + 2
        total_formula = "=H{}+H{}{}".format(r_sum, r_sum + 1, tail_extra)
    else:
        if service:
            print("[Writer] {}：模板无服务费行，现场服务开关未生效".format(
                os.path.basename(out_path)))

        r_total = r_sum + 1
        total_formula = "=H{}{}".format(r_sum, tail_extra)

    cap_formula = str(P_TOTAL["cells"][3][0] or "").replace(
        "H{}".format(r_total_src), "H{}".format(r_total))
    _write_row(ws, r_total, P_TOTAL, {
        3: cap_formula, 8: total_formula,
    })
    ws.merge_cells(start_row=r_total, start_column=3, end_row=r_total, end_column=6)
    ws.merge_cells(start_row=r_total, start_column=8, end_row=r_total, end_column=9)
    row = r_total + 1

    # 报价说明/TERMS 头 + 备注(行数随模板)
    for r in range(r_total_src + 1, max(protos) + 1):
        a_val = protos[r]["cells"].get(1, (None, None))[0]
        b_val = protos[r]["cells"].get(2, (None, None))[0]

        if a_val is None and b_val is None:
            continue

        _write_row(ws, row, protos[r], {1: a_val, 2: b_val})

        for c1, c2 in merge_by_row.get(r, []):
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)

        row += 1

    wb.properties.creator = "Z.ai"
    wb.save(out_path)


_WRITERS = {
    "COPA B": (_write_single, _write_multi),
    "COPA A": (_write_aw_single, _write_aw_multi),
}

# 平台秤/台秤业务的报价内容块模板(自真实单据提取，仅含报价内容单元格)
_SCALE_BLOCK_TPL = {
    "platform": "平台秤报价内容模板.xlsx",
    "bench": "台秤报价内容模板.xlsx",
}

# 单项目报价的品牌底板(公司抬头 + 报价内容： 标签所在区)
_SCALE_BASE_TPL = {
    "COPA B": "COPA B单项目报价模板.xlsx",
    "COPA A": "COPA A单项目报价模板.xlsx",
}


def _capture_block(block_ws, max_col=8):
    """内容块模板 -> 行快照列表[{height, cells:{col:(值,样式)}, merges:[(c1,c2)]}]"""
    rows = []

    for r in range(1, block_ws.max_row + 1):
        cells = {}

        for c in range(1, max_col + 1):
            cell = block_ws.cell(row=r, column=c)
            cells[c] = (cell.value, _style_of(cell))

        rows.append({
            "height": block_ws.row_dimensions[r].height,
            "cells": cells,
            "merges": [(m.min_col, m.max_col) for m in block_ws.merged_cells.ranges
                       if m.min_row == r],
        })

    return rows


def _write_scale_single(items, out_path, brand, service=False, lang="zh"):
    """平台秤/台秤单项目报价：品牌底板(公司抬头) + 业务内容块模板拼装。
    底板表头行(序号/No.)以下整体清空，内容块(整机行/服务费/价税合计/备注)按
    块内相对位置重写到表头行起，公式引用按目标行号重生成。"""
    item = items[0]
    brand_base = _tpl_name(_SCALE_BASE_TPL[brand], lang)
    block_name = _tpl_name(_SCALE_BLOCK_TPL[item["type"]], lang)

    base = load_workbook(os.path.join(TEMPLATE_DIR, brand_base))
    ws = base.worksheets[0]
    _stamp_date(ws)
    _stamp_company_contact(ws, items)

    # 底板表头行(A列='序号'/'No.')
    header_row = next(r for r in range(1, 16)
                      if ws.cell(row=r, column=1).value in ("序号", "No."))

    # 清空表头行起(含表头行的合并，块表头为独立单元格)的值/合并，
    # 底板原明细区由内容块整体接管
    for m in [m for m in list(ws.merged_cells.ranges) if m.min_row >= header_row]:
        ws.unmerge_cells(str(m))

    for r in range(header_row, ws.max_row + 1):
        for c in range(1, 10):
            ws.cell(row=r, column=c).value = None

    # 内容块搬运(现场服务关 -> 块内服务费行跳过，后续行上移一行)
    block = load_workbook(os.path.join(TEMPLATE_DIR, block_name)).worksheets[0]
    block_rows = _capture_block(block)

    def _is_total_row(snap):
        text = str(snap["cells"][1][0] or "")
        return text.startswith(("价税合计", "Total Amount", "Grand Total",
                                "Total:"))

    def _is_serve_row(snap):
        text = str(snap["cells"].get(2, (None,))[0] or "")
        return "服务费" in text or "Service Fee" in text

    total_src = next(i + 1 for i, row in enumerate(block_rows)
                     if _is_total_row(row))
    serve_src = next((i + 1 for i, row in enumerate(block_rows[:total_src - 1])
                      if _is_serve_row(row)), None)

    target = header_row

    for i, snap in enumerate(block_rows):
        if (not service) and serve_src and (i + 1 == serve_src):
            continue          # 服务费行整行跳过(后续行自动上移)

        r = target
        ws.row_dimensions[r].height = snap["height"]

        for c, (value, style) in snap["cells"].items():
            cell = cast(Any, ws.cell(row=r, column=c))
            cell.value = value
            _apply_style(cell, style)

        for c1, c2 in snap["merges"]:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)

        target += 1

    # 填整机行(块内第2行)：名称/描述/单位/数量；序号与总价公式按目标行重生成
    r_detail = header_row + 1
    name, model, unit, qty, desc = build_item_rows(item, lang=lang)[0][0]
    _set_cell(ws, r_detail, 2, name)
    _set_cell(ws, r_detail, 3, desc)
    _set_cell(ws, r_detail, 4, unit)
    _set_cell(ws, r_detail, 5, qty)
    _set_cell(ws, r_detail, 1, "=ROW()-{}".format(header_row))
    _set_cell(ws, r_detail, 7, "=E{0}*F{0}".format(r_detail))

    r_serve = (header_row + serve_src - 1) if (serve_src and service) else None

    if r_serve:
        _set_cell(ws, r_serve, 1, "=ROW()-{}".format(header_row))
        _set_cell(ws, r_serve, 7, "=E{0}*F{0}".format(r_serve))

    # 价税合计：SUM范围/大写公式引用按目标行号改写(跳过服务费行时合计行上移)
    r_total = header_row + total_src - 1 - (1 if (serve_src and not service) else 0)
    _set_cell(ws, r_total, 7, "=SUM(G{}:G{})".format(
        r_detail, r_serve or r_detail))
    cap = str(ws.cell(row=r_total, column=3).value or "")

    if cap:
        _set_cell(ws, r_total, 3, re.sub(
            r"G\d+", "G{}".format(r_total), cap))

    base.properties.creator = "Z.ai"
    base.save(out_path)


def _use_multi_template(items):
    """多项目模板判定：quote_id 最大>1；或模块业务选用多通道仪表——
    即便单项目报价，多通道仪表也需独立分项段输出，强制走多项目模板"""
    if max(p["quote_id"] for p in items) > 1:
        return True

    if len(items) == 1 and items[0].get("type") == "module":
        info = lookup_instrument(items[0].get("controller"))

        if info is not None and info["channels"] > 1:
            return True

    return False


def suggest_path(items, brand, lang="zh"):
    """按项目数与抬头生成默认保存路径(含时间戳)，供保存对话框预填；
    英文输出在文件名追加 _英文 标记"""
    multi = _use_multi_template(items)
    name = "{}报价单_{}项目_{}{}.xlsx".format(
        brand, "多" if multi else "单", datetime.now().strftime("%Y%m%d_%H%M%S"),
        "_英文" if lang == "en" else "",
    )
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    return os.path.join(OUTPUT_DIR, name)


def build_quote(items, brand="COPA B", out_path=None, out_dir=None,
                service=False, lang="zh"):
    """入口：items 为 Checklist 采集的结构化条目列表，brand 为报价抬头，
    lang 为输出语言("zh"中文/"en"英文，英文走 *_英文版 模板与英文文案)。
    quote_id 最大值 > 1 判为多项目报价。out_path 指定保存文件(保存对话框确认值)，
    缺省时按 out_dir(默认 output/) 自动命名。返回输出文件路径。"""
    if not items:
        raise ValueError("无可输出的条目")

    if brand not in _WRITERS:
        raise ValueError("不支持的报价抬头：{}(可选：{})".format(brand, "、".join(_WRITERS)))

    if lang not in _TEXT:
        raise ValueError("不支持的输出语言：{}(可选：zh/en)".format(lang))

    if out_path is None:
        out_path = suggest_path(items, brand, lang)
    else:
        parent = os.path.dirname(out_path)

        if parent:
            os.makedirs(parent, exist_ok=True)

    # 单项目且为秤业务 -> 整机单行制(真实单据格式，品牌底板+内容块模板)
    if (len(items) == 1 and max(p["quote_id"] for p in items) <= 1
            and items[0].get("type") in _SCALE_BLOCK_TPL):
        _write_scale_single(items, out_path, brand, service, lang)
        return out_path

    _WRITERS[brand][1 if _use_multi_template(items) else 0](
        items, out_path, service, lang)
    return out_path
